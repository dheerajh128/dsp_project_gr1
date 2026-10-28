# -*- coding: utf-8 -*-
"""Build the HEARTLINE DSP project report (DOCX)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Emu
from docx.enum.section import WD_ORIENT

ROOT = Path(r"D:\dsp_project")
OUT = ROOT / "report_final"
ASSETS = OUT / "assets"
LOGO = ASSETS / "mnnit_logo.png"
QR = ASSETS / "colab_qr.png"
BLOCK = ASSETS / "report_p3_img0.png"
FILTERS = ASSETS / "report_p6_img0.png"

NAVY = "1F4E79"
BLACK = "000000"

FIGS = []  # inventory rows
FIG_N = 0
TBL_N = 0


def _set_run_font(run, name="Calibri", size=12, bold=False, italic=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def _p_fmt(p, before=0, after=8, line=1.08, align="justify", keep_next=False):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.keep_with_next = keep_next
    if align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    elif align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT


def add_text(p, text, **kw):
    r = p.add_run(text)
    _set_run_font(r, **kw)
    return r


def shade_cell(cell, hex_color):
    tc = cell._tePr if hasattr(cell, "_tePr") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "BFBFBF")
        tcBorders.append(el)
    tcPr.append(tcBorders)


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cs = OxmlElement("w:cantSplit")
    trPr.append(cs)


def add_page_field(paragraph):
    run = paragraph.add_run()
    _set_run_font(run, size=9, color="333333")
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)


def add_toc_field(paragraph, instr_text):
    run = paragraph.add_run()
    r = run._r
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    i = OxmlElement("w:instrText")
    i.set(qn("xml:space"), "preserve")
    i.text = instr_text
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    t = OxmlElement("w:t")
    t.text = "Right-click and update this field if page numbers are blank."
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r.append(b)
    r.append(i)
    r.append(s)
    r.append(t)
    r.append(e)


def hrule(paragraph, color="1F4E79", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), sz)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def header_footer(section, cover=False):
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.different_first_page_header_footer = False
    hp = section.header.paragraphs[0]
    hp.clear()
    if cover:
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(hp, "HEARTLINE", name="Calibri", size=10, italic=True, color=NAVY)
        hrule(hp, NAVY, "12")
    else:
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_text(hp, "HEARTLINE  ·  Intelligent ECG Noise Classification", name="Calibri", size=9, italic=True, color="444444")
        hrule(hp, NAVY, "8")
    fp = section.footer.paragraphs[0]
    fp.clear()
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    hrule_top = fp._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    top = OxmlElement("w:top")
    top.set(qn("w:val"), "single")
    top.set(qn("w:sz"), "8")
    top.set(qn("w:space"), "6")
    top.set(qn("w:color"), NAVY)
    pBdr.append(top)
    hrule_top.append(pBdr)
    add_text(fp, "Motilal Nehru National Institute of Technology  |  ECN15101  |  Page ", size=9, color="333333")
    add_page_field(fp)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.7)
    section.header_distance = Cm(0.6)
    section.footer_distance = Cm(0.5)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)


def style_doc(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(12)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.space_after = Pt(8)
    pf.line_spacing = 1.08
    for style_name, size in (("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)):
        st = doc.styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(NAVY)
        st.font.italic = False
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        st.paragraph_format.space_before = Pt(12 if style_name != "Heading 1" else 4)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.0
        st.paragraph_format.keep_with_next = True
    cap = doc.styles["Caption"]
    cap.font.name = "Calibri"
    cap.font.size = Pt(10)
    cap.font.italic = True
    cap.font.color.rgb = RGBColor.from_string("222222")


def H(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    return p


def P(doc, text, size=12, bold=False, italic=False, align="justify", before=0, after=8, space=False):
    p = doc.add_paragraph()
    _p_fmt(p, before=before, after=after, align=align)
    add_text(p, text, size=size, bold=bold, italic=italic)
    return p


def bullets(doc, items):
    for it in items:
        p = doc.add_paragraph()
        _p_fmt(p, before=0, after=3, align="left")
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        add_text(p, "•  " + it, size=12)


def equation(doc, text):
    p = doc.add_paragraph()
    _p_fmt(p, before=4, after=4, align="center")
    add_text(p, text, name="Cambria Math", size=12, italic=True)
    return p


def table(doc, headers, rows, col_widths=None, font=9):
    global TBL_N
    t = doc.add_table(rows=1 + len(rows), cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = True
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _p_fmt(p, before=2, after=2, align="center")
        add_text(p, h, size=font, bold=True, color="FFFFFF")
        shade_cell(cell, NAVY)
        set_cell_border(cell)
    prevent_row_split(t.rows[0])
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            _p_fmt(p, before=1, after=1, align="left")
            add_text(p, str(val), size=font)
            if i % 2 == 1:
                shade_cell(cell, "F4F7FB")
            set_cell_border(cell)
        prevent_row_split(t.rows[i + 1])
    if col_widths:
        for row in t.rows:
            for j, w in enumerate(col_widths):
                row.cells[j].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def caption(doc, kind, text, chapter, source, experiment, origin=""):
    global FIG_N, TBL_N
    if kind == "Figure":
        FIG_N += 1
        n = FIG_N
        FIGS.append({
            "number": f"Figure {n}",
            "source_path": str(source),
            "caption": text,
            "experiment_id": experiment,
            "chapter": chapter,
            "origin": origin,
        })
    else:
        TBL_N += 1
        n = TBL_N
    p = doc.add_paragraph()
    _p_fmt(p, before=2, after=10, align="center")
    p.style = doc.styles["Caption"]
    add_text(p, f"{kind} {n}. {text}", size=10, italic=True)
    return n


def figure(doc, path, text, chapter, experiment, origin, width=6.15):
    p = doc.add_paragraph()
    _p_fmt(p, before=6, after=2, align="center", keep_next=True)
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width))
    n = caption(doc, "Figure", text, chapter, path, experiment, origin)
    return n


def cover(doc):
    header_footer(doc.sections[0], cover=True)
    p = doc.add_paragraph()
    _p_fmt(p, before=0, after=0, align="center")
    add_text(p, "DEPARTMENT OF ELECTRONICS AND COMMUNICATION ENGINEERING", size=11, bold=True, color=NAVY)
    p = doc.add_paragraph()
    _p_fmt(p, before=2, after=2, align="center")
    add_text(p, "Motilal Nehru National Institute of Technology", size=12, bold=True)
    p = doc.add_paragraph()
    _p_fmt(p, before=8, after=0, align="center")
    add_text(p, "PROJECT REPORT", size=12, bold=True, color=NAVY)
    hrule(p, NAVY, "16")
    p = doc.add_paragraph()
    _p_fmt(p, before=8, after=0, align="center")
    add_text(p, "INTELLIGENT ECG NOISE CLASSIFICATION", size=18, bold=True)
    p = doc.add_paragraph()
    _p_fmt(p, before=0, after=2, align="center")
    add_text(p, "AND DYNAMIC DSP FILTERING", size=18, bold=True)
    p = doc.add_paragraph()
    _p_fmt(p, before=0, after=4, align="center")
    add_text(p, "HEARTLINE", size=12, italic=True, color=NAVY)
    p = doc.add_paragraph()
    _p_fmt(p, before=2, after=2, align="center")
    r = p.add_run()
    r.add_picture(str(LOGO), width=Inches(1.32))
    p = doc.add_paragraph()
    _p_fmt(p, before=2, after=0, align="center")
    add_text(p, "B.Tech. — Electronics and Communication Engineering", size=11)
    p = doc.add_paragraph()
    _p_fmt(p, before=0, after=0, align="center")
    add_text(p, "Course: Digital Signal Processing  [ECN15101]", size=11)
    p = doc.add_paragraph()
    _p_fmt(p, before=0, after=4, align="center")
    add_text(p, "Group 1", size=12, bold=True)

    cap_p = doc.add_paragraph()
    _p_fmt(cap_p, before=2, after=2, align="center")
    add_text(cap_p, "Submitted by", size=11, bold=True, color=NAVY)

    students = [
        ("1", "Ayush Kumar Jha", "20244050"),
        ("2", "Bayapureddy Dheerajh Reddy", "20244051"),
        ("3", "Bhaskar Jalan", "20244052"),
        ("4", "Bhawani Singh Shekhawat", "20244053"),
    ]
    t = doc.add_table(rows=1 + len(students), cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ("S. No.", "Name of Student", "Roll Number")
    for j, h in enumerate(headers):
        cell = t.rows[0].cells[j]
        cell.text = ""
        p = cell.paragraphs[0]
        _p_fmt(p, before=1, after=1, align="center")
        add_text(p, h, size=10, bold=True, color="FFFFFF")
        shade_cell(cell, NAVY)
        set_cell_border(cell)
    for i, (n, name, roll) in enumerate(students):
        for j, val in enumerate((n, name, roll)):
            cell = t.rows[i + 1].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            _p_fmt(p, before=1, after=1, align="center" if j != 1 else "left")
            add_text(p, val, size=10)
            set_cell_border(cell)
    for row in t.rows:
        row.cells[0].width = Cm(2.0)
        row.cells[1].width = Cm(9.5)
        row.cells[2].width = Cm(3.5)

    # QR and submission side by side
    spacer = doc.add_paragraph()
    _p_fmt(spacer, before=2, after=0, align="center")
    box = doc.add_table(rows=1, cols=2)
    box.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = box.rows[0].cells
    left.text = ""
    lp = left.paragraphs[0]
    _p_fmt(lp, before=2, after=0, align="center")
    rr = lp.add_run()
    rr.add_picture(str(QR), width=Inches(1.38))
    lp2 = left.add_paragraph()
    _p_fmt(lp2, before=1, after=0, align="center")
    add_text(lp2, "Google Colab notebook", size=9, italic=True, color="333333")
    right.text = ""
    rp = right.paragraphs[0]
    _p_fmt(rp, before=6, after=0, align="center")
    add_text(rp, "Submitted to", size=12, bold=True, color=NAVY)
    for line in ("Dr. Shishir Maheshwari", "Prof. V. Krishna Rao Kandanvli"):
        rp2 = right.add_paragraph()
        _p_fmt(rp2, before=2, after=0, align="center")
        add_text(rp2, line, size=12, bold=True)
    for cell in (left, right):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()
        tcBorders = OxmlElement("w:tcBorders")
        for edge in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "nil")
            tcBorders.append(el)
        tcPr.append(tcBorders)


def body_start(doc):
    new = doc.add_section()
    header_footer(new, cover=False)
    # continue page numbers
    sectPr = new._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is not None:
        sectPr.remove(pg)


def build():
    doc = Document()
    style_doc(doc)
    core = doc.core_properties
    core.title = "Intelligent ECG Noise Classification and Dynamic DSP Filtering"
    core.author = "Group 1"
    core.subject = "HEARTLINE — ECN15101 project report"
    core.category = "B.Tech. ECE project report"

    cover(doc)
    body_start(doc)

    # ---------- front matter ----------
    H(doc, "Abstract")
    P(doc, "Electrocardiogram (ECG) recordings are low-amplitude signals. Baseline wander (BW), muscle artifact (MA) and power-line interference (PLI) occupy different parts of the spectrum, so a single fixed filter is a poor match for every window. This project, HEARTLINE, separates noise recognition from restoration. Clean ECG from the MIT-BIH Arrhythmia Database, sampled at 360 Hz, is mixed with one noise class at a time. Real BW and MA come from the MIT-BIH Noise Stress Test Database. PLI is a synthetic 50 Hz sine, scaled so that the mixture meets a chosen signal-to-noise ratio (SNR).")
    P(doc, "The classical path extracts hand-crafted time and spectral features and classifies the window as BW, MA or PLI. On the largest completed classical experiment — 9,000 windows of 4 s, 30 records, a record-level split and SNR values of 0, 5, 10 and 15 dB — a histogram gradient boosting classifier reached a held-out accuracy of 0.8005. A tuned random forest on 4,500 windows of 2 s reached 0.7766. In every three-class confusion matrix, PLI is essentially solved and the remaining errors are exchanges between BW and MA. Accuracy falls as the SNR rises, because milder noise is harder to recognise.")
    P(doc, "A second path converts each window to a log-power spectrogram and classifies the image. Saved evaluations give a test accuracy of 0.8772 for the six-hour CNN run, 0.8765 for ResNet-101 on the same test size (1,344 windows), and 0.8590 for ResNet-101 on a 2,256-window test set. The current source implements a ResNet-50 transfer-learning pipeline, but that run did not leave an evaluation file in the repository. Figures quoted only in an earlier seven-page project document (89.67%, 90.87% and 89.80%) are reported as document claims, not as files that can be reopened here.")
    P(doc, "Class-routed filtering is specified in that earlier document: a 4th-order Butterworth high-pass at 0.5 Hz for BW, a 4th-order Butterworth low-pass at 35 Hz for MA, and a 50 Hz IIR notch with Q = 30 for PLI. No filter function is present in the current Python sources, and no SNR-improvement table was saved. The work is a DSP and machine-learning study on public data. It is not a clinical device and it has not been deployed.")

    H(doc, "Index")
    P(doc, "The entries below are a Word table of contents. Page numbers are filled when fields are updated.", size=11, italic=True, align="left")
    tp = doc.add_paragraph()
    add_toc_field(tp, ' TOC \\o "1-2" \\h \\z \\u ')

    H(doc, "List of Figures")
    for item in [
        "Figure 1. Clean ECG, record 100 (first 10 s).",
        "Figure 2. Electrode-motion ECG, record 118e_6 (first 10 s).",
        "Figure 3. Clean and noisy 4 s windows.",
        "Figure 4. Windowing: 4 s at 360 Hz is 1,440 samples.",
        "Figure 5. One exploratory six-feature vector.",
        "Figure 6. Exploratory feature distributions, Clean versus Noisy.",
        "Figure 7. Classification-driven filter block diagram (from the earlier report).",
        "Figure 8. Improved 2 s random forest, confusion matrix.",
        "Figure 9. Improved 2 s feature importance.",
        "Figure 10. Improved 2 s accuracy by SNR.",
        "Figure 11. Final 4 s confusion matrix.",
        "Figure 12. Final 4 s accuracy by SNR.",
        "Figure 13. Final 4 s feature-importance figure.",
        "Figure 14. Exploratory two-class confusion matrix.",
        "Figure 15. Exploratory feature importance.",
        "Figure 16. CNN six-hour confusion matrix.",
        "Figure 17. ResNet-101 six-hour confusion matrix.",
        "Figure 18. ResNet-101 ten-hour confusion matrix.",
        "Figure 19. Filter examples reproduced from the earlier report.",
    ]:
        P(doc, item, size=11, align="left", before=0, after=2)

    H(doc, "List of Tables")
    for item in [
        "Table 1. Abbreviations.",
        "Table 2. Six core features.",
        "Table 3. Class-to-filter map (specified, not coded).",
        "Table 4. Repository locations.",
        "Table 5. Classical three-class results.",
        "Table 6. Improved 2 s per-class scores.",
        "Table 7. Final 4 s per-class scores.",
        "Table 8. Saved CNN and ResNet-101 scores.",
        "Table 9. Scores found only in the earlier PDF.",
        "Table 10. Exploratory feature means.",
    ]:
        P(doc, item, size=11, align="left", before=0, after=2)

    H(doc, "Abbreviations")
    table(doc,
          ["Abbreviation", "Meaning"],
          [
              ["BW", "Baseline wander"],
              ["MA", "Muscle artifact (electromyographic noise)"],
              ["PLI", "Power-line interference"],
              ["EM", "Electrode-motion noise (used only in the exploratory figures)"],
              ["ECG", "Electrocardiogram"],
              ["SNR", "Signal-to-noise ratio"],
              ["STFT", "Short-time Fourier transform"],
              ["RF", "Random forest"],
              ["HGB", "Histogram-based gradient boosting"],
              ["CNN", "Convolutional neural network"],
              ["CRNN", "Convolutional recurrent neural network (not present in this repository)"],
              ["DSP", "Digital signal processing"],
              ["IIR", "Infinite impulse response"],
              ["PSD", "Power spectral density"],
              ["MIT-BIH", "Massachusetts Institute of Technology–Beth Israel Hospital databases"],
              ["NSTDB", "MIT-BIH Noise Stress Test Database"],
              ["MLII", "Modified limb lead II"],
          ],
          font=10)
    caption(doc, "Table", "Abbreviations used in this report.", "Abbreviations", "", "nomenclature", "nomenclature")

    # ---------- 1 ----------
    H(doc, "1. Introduction")
    H(doc, "1.1 Background", 2)
    P(doc, "A diagnostic ECG is a small voltage measured on the body surface. At the 360 Hz sampling rate of the MIT-BIH Arrhythmia Database, a few seconds of recording already contain the P wave, QRS complex and T wave that later analysis depends on. The same recording also collects disturbances that are not cardiac. Three of them are treated in this project because they are both common and spectrally different.")
    P(doc, "Baseline wander is a slow drift, typically well below 0.5 Hz, caused by respiration, poor electrode contact and body movement. Muscle artifact is broadband activity from skeletal muscle. It overlaps the band that also contains the QRS. Power-line interference is a narrow tone at the mains frequency. In this project that frequency is 50 Hz, which is the mains frequency assumed by the synthesised interferer. An inappropriate filter does real harm: a high-pass cutoff that is too high distorts the ST segment, and a low-pass cutoff that is too low rounds the QRS.")
    P(doc, "The engineering problem is therefore not only “remove noise”. It is “decide which noise is dominant in this window, then apply the filter that matches that decision”. Recognition and restoration are different operations. Recognition can be statistical. Restoration, in the design used here, stays with fixed filters whose frequency response can be written down.")

    H(doc, "1.2 Motivation", 2)
    P(doc, "A fixed filter applied to every window is simple and often wrong. Manual choice of the filter does not scale to long recordings and is not repeatable. The first method tried in the project was a random forest on a small set of scalar features: RMS, variance, kurtosis, low-frequency power, power near 50 Hz and high-frequency power. Those features are interpretable, and the 50 Hz band in particular separates PLI. They do not represent how energy is arranged jointly in time and frequency, and BW and MA remain easy to confuse once the evaluation is done on patients who were not seen in training.")
    P(doc, "The later line of work therefore builds a time–frequency image of each window and asks a convolutional network to name the noise. The network output is intended to be the only control signal. The filters themselves do not change with the patient. That split is the design stated in the earlier project write-up and drawn in its block diagram. The repository implements the data, the classical classifiers and the spectrogram network. It does not implement the filter functions.")

    H(doc, "1.3 Objectives", 2)
    bullets(doc, [
        "Build labelled noisy-ECG windows from MIT-BIH clean records and NSTDB noise, with one noise class per window and a controlled SNR.",
        "Compare window lengths of 2 s and 4 s, and compare a small single-record split with record-level splits.",
        "Classify BW, MA and PLI from hand-crafted features, and record accuracy, per-class scores, confusion matrices and behaviour across SNR.",
        "Train spectrogram classifiers (a CNN and ResNet models) and compare them with the classical results on the numbers that were actually saved.",
        "Specify, and where a figure already exists, illustrate a class-to-filter map for BW, MA and PLI.",
    ])

    H(doc, "1.4 Scope", 2)
    P(doc, "The study uses public PhysioNet records. It does not collect new patient data. Noise classes are injected one at a time, so the label is the class that was added, not a clinician’s reading of a real mixture. Overlapping BW and MA in the same window is outside the experiments. No arrhythmia detector is run on the restored signal. No deployment, embedded implementation or clinical validation is claimed.")
    P(doc, "Two result groups must be kept apart. Group A is the original three-class BW/MA/PLI work under outputs/milestone3_results/. Group B is an exploratory clean-versus-noisy demonstration in random_forest_figures/, built from record 100 and the electrode-motion stress record 118e_6. Group B explains waveforms, windowing and features. Its perfect test score is not a BW/MA/PLI result.")

    # ---------- 2 ----------
    H(doc, "2. Literature Review and Theoretical Background")
    H(doc, "2.1 ECG noise and signal quality", 2)
    P(doc, "Satija, Ramkumar and Manikandan review ECG signal-quality assessment and treat noise identification as a step that should precede analysis [6]. Li, Rajagopalan and Clifford use engineered quality features with classical machine learning for the same purpose [7]. Both support the project’s first stage: a compact feature vector and a classical classifier. They also show why a single quality score is not enough here. The required output is a noise class, because the class chooses the filter.")
    P(doc, "The clean signals are the MIT-BIH Arrhythmia Database [11], [12], accessed through PhysioNet. The noise records bw and ma, and the stress-test record 118e_6 used only for illustration, are from the MIT-BIH Noise Stress Test Database [10], [12]. The database note distributed with the local copy states that 118e_6 is record 118 with electrode-motion noise at an SNR of −6 dB. That is a different noise from BW, MA and PLI.")

    H(doc, "2.2 Why the spectrum decides the filter", 2)
    P(doc, "BW energy sits near dc. A high-pass Butterworth response is the standard linear tool for that drift, and Romero and colleagues discuss baseline-wander removal in that setting [9]. PLI is a line. A narrow IIR notch can reject it, but a fixed notch is sensitive to frequency error and can ring; Chen and colleagues discuss those limits [8]. MA is the difficult case. It is broadband and overlaps the ECG, so a low-pass filter can only suppress the part of the artifact that lies above the cutoff. Energy inside the QRS band remains.")
    P(doc, "The literature also argues that suppression should follow the noise scenario rather than one permanent filter [5]. Neural selection of a denoising filter is reported by Pravin and Ojha [1]. A one-dimensional CNN has been used to detect wandering baseline and mains interference [2]. Spectrogram images as CNN input follow the same time–frequency idea as [4]. Transfer learning from a large image model is discussed for ECG in [3]. This project uses an ImageNet-pretrained ResNet-50 in source code for that reason. The architecture was not designed by the group.")

    H(doc, "2.3 What is used, and what is not claimed", 2)
    P(doc, "The citations above are the references carried in Complete_Project_Report.pdf. Several of those entries are incomplete in the source document (missing full titles, initials or venues). They are reproduced in the reference list in the form available there, and they are not silently completed. No additional papers were invented to fill the gaps. Random forests and residual networks are described from the behaviour of the code and the saved models, not from a new bibliographic claim.")

    # ---------- 3 ----------
    H(doc, "3. Methodology")
    H(doc, "3.1 Records and sampling", 2)
    P(doc, "Records are read with the WFDB library (src/data_loader.py). Discovery is by header file, not by a hard-coded list. Every experiment that states a sampling rate uses fs = 360 Hz. A header check on record 100 shows two channels, MLII and V5, with MLII stored first. The single-window generator in src/data_generator.py prefers a channel named MLII and falls back to channel 0. The multi-record generator in src/experiment_2s.py always takes channel 0. For record 100 those choices coincide. They need not coincide on every MIT-BIH record, and the report does not claim that every window is MLII.")
    P(doc, "Noise record bw supplies baseline wander and noise record ma supplies muscle artifact, each from channel 0 of the NSTDB file. Both headers in the local copy show fs = 360 Hz and 650,000 samples, about 30 minutes. PLI is not read from a file. src/data_generator.py synthesises")
    equation(doc, "w_PLI[n] = sin(2 π · 50 · n / fs)")
    P(doc, "with unit amplitude. That amplitude is only a shape. The SNR stage rescales whatever noise vector it is given.")

    H(doc, "3.2 One noise class at a controlled SNR", 2)
    P(doc, "Each window contains exactly one added disturbance. The scale α is chosen so that the power ratio matches a target SNR S in decibels:")
    equation(doc, "SNR = 10 log10( Px / Pw ),    Px = mean(x²),    Pw = mean(w²)")
    equation(doc, "α = sqrt( (Px / 10^(S/10)) / Pw ),    x_noisy = x + α w")
    P(doc, "This is the implementation in calculate_noise_scaling and add_noise_at_snr. A near-zero noise segment is rejected and another start index is drawn. The earlier project document left open whether PLI was real or synthetic. The source settles it: PLI is synthetic. BW and MA are real NSTDB segments.")
    P(doc, "The SNR grids that were actually stored are not all the same. The 180-window and 900-window sets use 0, 5 and 10 dB. The 4,500-window and 9,000-window sets use 0, 5, 10 and 15 dB. A higher SNR is a cleaner window. It is a harder classification problem, not an easier one.")

    H(doc, "3.3 Windows and splits", 2)
    P(doc, "A 4 s window at 360 Hz contains 1,440 samples. A 2 s window contains 720 samples. Starts are random. Windows may overlap, and the same region of a record can be drawn more than once. Statements such as “10 hours” in main.py mean the sum of window lengths (9,000 × 4 s = 36,000 s), not 10 hours of non-overlapping ECG.")
    P(doc, "The earliest 4 s set has 180 windows and no record identifier in its metadata. Its summary describes a single-record experiment with a random 126/54 split, which is the 70/30 split of the original classifier. Later experiments split by record: no ECG record contributes windows to both sides. The long BW and MA recordings are cut in time at the same fraction, so a noise segment used in training is not reused in testing. PLI is regenerated for every window, so it has no recorded start index.")
    P(doc, "There is a small integer effect in the multi-record generator. The per-SNR target is a rounded share of the per-class target, so the saved test count is 1,128 or 2,256 rather than the round number written in a comment. The metadata files are the counts used in this report.")

    figure(doc, ROOT / "random_forest_figures" / "100_raw_ecg.png",
           "First 10 s of MIT-BIH record 100, channel 0 (MLII). The QRS complexes are sharp and the baseline is nearly flat. This trace is an illustration of a clean input. It is not a BW, MA or PLI test sample.",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.2)
    figure(doc, ROOT / "random_forest_figures" / "118e_6_raw_ecg.png",
           "First 10 s of NSTDB record 118e_6. The database note identifies this file as record 118 plus electrode-motion noise at −6 dB. The baseline swings by about 1 mV. This is not a label in the three-class experiments.",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.2)
    figure(doc, ROOT / "random_forest_figures" / "clean_vs_noisy.png",
           "The first 4 s window of record 100 (top) and of record 118e_6 (bottom), plotted from the exploratory script. Electrode motion changes both the baseline and the apparent QRS amplitude.",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.2)
    figure(doc, ROOT / "random_forest_figures" / "windowing_example.png",
           "Windowing on record 100. Four seconds at 360 Hz is 1,440 samples, which is the window length of the final classical experiment and of the ResNet-50 source.",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.2)

    H(doc, "3.4 Hand-crafted features", 2)
    P(doc, "The six features named in the earlier write-up are implemented in src/feature_extraction.py. RMS is the square root of the mean square of the window and therefore includes any dc offset. Variance is the population variance about the mean. Kurtosis is scipy.stats.kurtosis on the raw window (Fisher definition, biased estimator). Spectral features are computed on a zero-mean copy. Welch’s method uses a segment length equal to the window, so the frequency spacing is 0.25 Hz for 1,440 samples and 0.5 Hz for 720 samples.")
    table(doc,
          ["Feature", "Definition in the three-class code", "Why it can separate classes"],
          [
              ["RMS", "sqrt(mean(x²))", "Overall energy. Sensitive to large wander and to strong artifact."],
              ["Variance", "population variance of x", "Spread about the mean, without the dc term that RMS keeps."],
              ["Kurtosis", "Fisher kurtosis of x", "Peakedness. A clean QRS is heavy-tailed; broadband noise fills the distribution."],
              ["Low-frequency power", "Welch power from 0.05 Hz to 0.5 Hz", "Intended to track baseline wander."],
              ["Power near 50 Hz", "Welch power from 49 Hz to 51 Hz", "A narrow measure of mains interference."],
              ["High-frequency power", "Welch power from 20 Hz to Nyquist", "Intended to track muscle artifact."],
          ],
          font=9)
    caption(doc, "Table", "Six core features used by the three-class pipeline. Ratios of low-frequency and high-frequency power to total power are also returned.",
            "3. Methodology", "feature_extraction.py", "implementation")
    P(doc, "Two later experiments add features. The saved 2 s improved forest has ten inputs. Its importance plot names them: the six core features except the two power ratios, plus zero-crossing rate, spectral entropy, dominant frequency and the low-to-high band-power ratio. The 4 s model file has fourteen inputs, which matches the current extended set: the eight values from feature_extraction.py plus zero-crossing rate, spectral entropy, dominant frequency, low-to-high ratio, mean absolute successive difference and Hjorth mobility.")
    P(doc, "The exploratory script in random_forest_visualizations.py is a different definition and must not be quoted as the three-class feature code. It removes the mean before RMS, uses Welch segments of at most 1,024 samples, and integrates 0.1–5 Hz, 48–52 Hz and 20–100 Hz. Those plots are still useful, because they show the shape of a feature vector and the gap between record 100 and record 118e_6.")
    figure(doc, ROOT / "random_forest_figures" / "example_feature_vector.png",
           "Six exploratory features from one window of record 100. Bar height is not importance: the features do not share a unit. Bands differ from the three-class extractor.",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.0)
    figure(doc, ROOT / "random_forest_figures" / "feature_distributions.png",
           "Distributions of the six exploratory features on 451 windows from each record. Means from the saved CSV: clean RMS 0.189 and variance 0.036, against noisy RMS 2.021 and variance 8.073. Low-frequency power (0.1–5 Hz) averages 0.0059 (clean) and 5.03 (noisy). Kurtosis averages 26.7 (clean) and 2.76 (noisy).",
           "3. Methodology", "exploratory_group_B", "exploratory_visualization", width=6.15)

    H(doc, "3.5 Classical classifiers", 2)
    P(doc, "src/classifier.py trains a random forest (100 trees, random_state 42) and evaluates accuracy, a classification report and a confusion matrix on a stratified 30% hold-out. That is the setup consistent with the 180-window experiment. Later scripts change the forest size, add a record-level split, and compare the forest with histogram gradient boosting. The saved 4 s model is HistGradientBoostingClassifier with max_iter = 300, learning_rate = 0.05 and random_state = 42, which is the constructor in main.py. Its file name still says random_forest. The loaded object, and the text file that says “Model: Gradient Boosting”, are the evidence used here. The forest’s own test accuracy in that showdown was printed by the script and was not written into the report file, so it is not stated.")
    P(doc, "The improved 2 s forest was tuned by randomised search. The saved parameters are 300 trees, max_depth 30, min_samples_split 2, min_samples_leaf 4 and max_features sqrt. The search space in the current experiment_improved.py does not contain depth 30 or a leaf size of 4, and it ignores the n_iter argument passed by main.py. The saved model is therefore from an earlier version of the search. The result file, not the current search grid, is quoted.")

    H(doc, "3.6 Spectrograms and the ResNet-50 source", 2)
    P(doc, "src/cnn_pipeline.py and main_cnn.py convert each 1,440-sample window with scipy.signal.spectrogram at nperseg = 64 and noverlap = 48. The hop is 16 samples. For a 4 s window the grid before resizing is 33 frequency bins by 87 time frames. No window argument is passed, so SciPy’s default Tukey window is used. Power is converted by 10 log10(Sxx + 1e−10), scaled to 0–255 inside each image, copied to three channels, resized to 128 × 128, and passed through the ResNet-50 preprocess function. No spectrogram image was saved, so these dimensions are the consequence of the call, not a measurement from a stored picture.")
    P(doc, "The network is ResNet-50 with ImageNet weights and without its original head. The new head is global average pooling, batch normalisation, dropout 0.4, a 128-unit ReLU dense layer with L2 factor 1e−4, batch normalisation, dropout 0.4, and a 3-way softmax with the same L2 factor. Training is written as 5 epochs with the backbone frozen (Adam, learning rate 1e−3, batch 64) and then 15 epochs with the last 25 base-model layers trainable (Adam, 1e−5). The loss is categorical cross-entropy with label smoothing 0.1. The validation set in this script is the patient hold-out that is also used as the test set, so checkpointing looks at the test patients. The script writes to outputs/milestone3_results/resnet50_opt2_opt3. That directory is not in the repository. The run is implemented and not evidenced by a saved score.")
    P(doc, "The earlier project document describes a different head: a 512-unit layer, dropout 0.5, and 45 fine-tuning epochs, and it quotes test accuracies of 89.67%, 90.87% and 89.80% on 4,500 windows. Those settings and scores are not in the result files. Both descriptions are kept, in different sections, so that a reader can see the conflict.")

    H(doc, "3.7 Intended filter routing", 2)
    P(doc, "No Butterworth or notch function exists under src/. Classifier comments say that milestone 3 does not filter. The routing below is the design in the earlier project document, including the order, cutoffs and Q that the document states. Forward–backward filtering is marked there as unconfirmed, and it is unconfirmed here as well.")
    table(doc,
          ["Predicted class", "Filter stated in the earlier report", "Role"],
          [
              ["BW", "4th-order Butterworth high-pass, cutoff 0.5 Hz", "Attenuate slow drift while keeping ST content above the cutoff."],
              ["MA", "4th-order Butterworth low-pass, cutoff 35 Hz", "Attenuate high-frequency EMG. Artifact below 35 Hz remains."],
              ["PLI", "IIR notch at 50 Hz, Q = 30", "Narrow rejection. The document estimates the −3 dB width as f0/Q ≈ 1.7 Hz."],
          ],
          font=9)
    caption(doc, "Table", "Class-to-filter map as specified in Complete_Project_Report.pdf. These filters are not implemented in the current source tree.",
            "3. Methodology", "Complete_Project_Report.pdf", "prior_report_not_in_source")
    figure(doc, BLOCK,
           "Block diagram reproduced from Complete_Project_Report.pdf. The drawn path is spectrogram, ResNet-50, then one of three fixed filters. The repository also contains the classical feature path, which this drawing does not show.",
           "3. Methodology", "prior_project_report", "prior_report_figure", width=6.3)

    H(doc, "3.8 Metrics", 2)
    P(doc, "Accuracy is the fraction of test windows whose predicted label equals the injected class. Because every saved three-class test set is balanced, balanced accuracy equals accuracy in these files, and that equality is not a separate finding. Precision, recall and F1 are the usual per-class scores. The confusion-matrix rows are true labels and the columns are predictions, in the order BW, MA, PLI unless a figure says otherwise. SNR curves are accuracies on the test windows of one SNR, with the model left unchanged.")

    # ---------- 4 ----------
    H(doc, "4. Implementation")
    H(doc, "4.1 Code and saved objects", 2)
    P(doc, "The Python package under src/ is the implementation. data_loader.py reads WFDB records. data_generator.py mixes noise at a target SNR. feature_extraction.py and experiment_improved.py build features. classifier.py, experiment_2s.py and experiment_improved.py train and plot. main.py is the 4 s, 30-record, gradient-boosting comparison. main_cnn.py and cnn_pipeline.py are the ResNet-50 spectrogram pipeline. random_forest_visualizations.py writes only the exploratory figures. Training uses NumPy, SciPy, scikit-learn, Matplotlib, joblib and, for the network, TensorFlow and Keras. No requirements file pins the versions.")
    table(doc,
          ["Path", "What it is"],
          [
              ["data/mit_db, data/noise_stress", "Local MIT-BIH and NSTDB files, including bw, ma and 118e_6."],
              ["outputs/noisy_dataset", "180 windows, 1,440 samples, SNR 0/5/10. Single-record 4 s set."],
              ["outputs/noisy_dataset_2s", "900 windows, 720 samples, 15 records, SNR 0/5/10."],
              ["outputs/noisy_dataset_2s_improved", "4,500 windows, 720 samples, SNR 0/5/10/15."],
              ["outputs/noisy_dataset_4s_final", "9,000 windows, 1,440 samples, 30 records, SNR 0/5/10/15."],
              ["models/*.joblib", "Four saved classical models. The 4 s file is gradient boosting."],
              ["outputs/milestone3_results/", "Metrics and figures for the runs that were kept."],
              ["random_forest_figures/", "Exploratory clean-versus-noisy plots and a two-class score."],
          ],
          font=9)
    caption(doc, "Table", "Repository locations that support the numerical claims.", "4. Implementation", "repository", "source_map")

    H(doc, "4.2 Experiments that were not found", 2)
    P(doc, "The directories crnn_6hours, resnet50_50ep and resnet50_opt2_opt3 are not in the tree. There is no CRNN source file. There is no saved weight file for ResNet-50 or ResNet-101. The folders cnn_6hours, resnet101_6hours and resnet101_10hours contain a text report and two confusion matrices each, and not a training log, an epoch count or a learning curve. Architecture details for those three runs are limited to the folder names and the label set BW/MA/PLI.")

    # ---------- 5 ----------
    H(doc, "5. Results and Analysis")
    H(doc, "5.1 Three-class classical experiments", 2)
    P(doc, "Table 5 collects every classical three-class score that is backed by a file in the repository. The two small experiments look stronger and are the weaker evidence: 54 and 180 test windows, and, in the 180-window case, windows drawn from one recording. The 4,500-window and 9,000-window results use held-out records and are the scores this report treats as the classical performance.")
    table(doc,
          ["Experiment", "Model", "Windows (test)", "SNR (dB)", "Test accuracy"],
          [
              ["Previous 4 s, one record", "Random forest, 100 trees, 6 features", "180 (54)", "0, 5, 10", "0.8333"],
              ["Previous 2 s, initial", "Random forest, 200 trees, 6 features", "900 (180)", "0, 5, 10", "0.8500"],
              ["Improved 2 s", "Random forest, 300 trees, depth 30, 10 features", "4,500 (1,128)", "0, 5, 10, 15", "0.7766"],
              ["Final 4 s", "Histogram gradient boosting, 300 iterations, 14 features", "9,000 (2,256)", "0, 5, 10, 15", "0.8005"],
          ],
          font=8)
    caption(doc, "Table", "Group A classical results. Accuracies are taken from the saved summaries. The 2 s initial and 4 s single-record models are identified by joblib file name and feature count, which the summary text does not repeat.",
            "5. Results", "milestone3_results", "original_experiment")

    P(doc, "Previous 4 s. The summary records 1,440 samples per window, 180 windows, a 126/54 split and accuracy 83.33%. The confusion matrix is [[14, 4, 0], [5, 13, 0], [0, 0, 18]] for labels BW, MA, PLI. All 18 PLI windows are correct. Four BW windows are called MA and five MA windows are called BW. The matching joblib object has 100 trees and six features. Impurity importances are 0.0966 (RMS), 0.0961 (variance), 0.0771 (kurtosis), 0 (low-frequency power), 0.4363 (50 Hz power) and 0.2938 (high-frequency power).")
    P(doc, "Previous 2 s. Nine hundred windows of 720 samples, 15 records (12 train, 3 test: 107, 122 and 209), SNR 0, 5 and 10 dB, test size 180, accuracy 85%. The matrix is [[43, 17, 0], [10, 50, 0], [0, 0, 60]]. The 2 s joblib file has 200 trees and the same six-feature importance pattern: 50 Hz power 0.497, high-frequency power 0.222, and low-frequency power again 0. At 2 s the Welch spacing is 0.5 Hz, so the 0.05–0.5 Hz band contains little more than a single bin. That is a credible reason for a dead feature. It is not proved by an ablation, because no ablation file was saved.")
    P(doc, "Improved 2 s. This is the experiment behind the “about 77.7%” sentence in the earlier write-up. The exact accuracy is 0.7766, with balanced accuracy 0.7766 and macro F1 0.7766. An untuned 100-tree baseline on the same split is recorded as 0.7571. The test set has 1,128 windows, 376 per class. Records in the summary are 15 for training and five for testing (105, 107, 122, 123, 209).")
    table(doc,
          ["Class", "Precision", "Recall", "F1", "Support"],
          [
              ["BW", "0.66", "0.67", "0.67", "376"],
              ["MA", "0.67", "0.66", "0.66", "376"],
              ["PLI", "1.00", "1.00", "1.00", "376"],
          ],
          font=10)
    caption(doc, "Table", "Improved 2 s random forest, per-class scores from classification_report.txt. Counts in the raw confusion matrix are BW 252 correct and 124 called MA; MA 128 called BW and 248 correct; PLI 376 correct.",
            "5. Results", "improved_2s", "original_experiment")
    figure(doc, ROOT / "outputs/milestone3_results/improved_2s/confusion_matrix_raw.png",
           "Improved 2 s random forest, raw confusion matrix on 1,128 held-out windows. Every error is between BW and MA.",
           "5. Results", "improved_2s", "original_experiment", width=4.8)
    figure(doc, ROOT / "outputs/milestone3_results/improved_2s/feature_importance.png",
           "Impurity importance of the saved 10-feature forest. The loaded values are 0.4534 for 50 Hz power, 0.1276 for dominant frequency, 0.1156 for high-frequency power, 0.0810 for zero-crossing rate, 0.0701 for spectral entropy, then variance 0.0533, RMS 0.0503 and kurtosis 0.0487. Low-frequency power and the low-to-high ratio are 0.",
           "5. Results", "improved_2s", "original_experiment", width=5.8)
    figure(doc, ROOT / "outputs/milestone3_results/improved_2s/accuracy_by_snr.png",
           "Improved 2 s test accuracy against SNR. The curve falls from 0 dB to 15 dB. Exact per-SNR values were not written beside the figure, so they are not quoted.",
           "5. Results", "improved_2s", "original_experiment", width=5.2)

    P(doc, "Final 4 s. Thirty records, listed in the dataset summary, produce 9,000 windows. Training uses 6,744 windows from 22 records. Testing uses 2,256 windows from records 104, 105, 107, 122, 209, 210, 228 and 232, with 752 windows per class. The selected model is histogram gradient boosting. Accuracy is 0.8005, balanced accuracy is 0.8005, macro F1 is 0.8001 and weighted F1 is 0.8001. Recomputing predictions from the saved model and the current 14-feature function on the stored test windows returns accuracy 0.80053, which rounds to the saved 0.8005. The same pass gives per-SNR accuracy 0.8617, 0.8333, 0.7766 and 0.7305 at 0, 5, 10 and 15 dB, with 564 windows at each SNR. Those four figures are a check of the saved model, not a new training run.")
    table(doc,
          ["Class", "Precision", "Recall", "F1", "Support"],
          [
              ["BW", "0.72", "0.66", "0.69", "752"],
              ["MA", "0.69", "0.74", "0.71", "752"],
              ["PLI", "1.00", "1.00", "1.00", "752"],
          ],
          font=10)
    caption(doc, "Table", "Final 4 s gradient boosting, per-class scores from classification_report.txt. Matrix counts: BW 496 correct, 254 called MA, 2 called PLI; MA 193 called BW and 559 correct; PLI 751 correct and 1 called BW.",
            "5. Results", "final_4s", "original_experiment")
    figure(doc, ROOT / "outputs/milestone3_results/final_4s/confusion_matrix_raw.png",
           "Final 4 s model, raw confusion matrix on 2,256 test windows. PLI is separated. BW and MA still exchange about one third of their windows.",
           "5. Results", "final_4s", "original_experiment", width=4.8)
    figure(doc, ROOT / "outputs/milestone3_results/final_4s/snr_performance.png",
           "Final 4 s test accuracy against SNR. The recomputed values are 0.8617 (0 dB), 0.8333 (5 dB), 0.7766 (10 dB) and 0.7305 (15 dB).",
           "5. Results", "final_4s", "original_experiment", width=5.2)
    figure(doc, ROOT / "outputs/milestone3_results/final_4s/feature_importance.png",
           "Feature-importance figure stored with the 4 s results. The title says random forest and the axes list the 14 extended features, led by zero-crossing rate and 50 Hz power. The saved champion is gradient boosting, which does not store this impurity vector. The figure is the random-forest view of those features, not an explanation of the selected booster.",
           "5. Results", "final_4s", "original_experiment", width=5.8)
    P(doc, "Reading the four classical results together: moving from a single record and a short test to a record-level test lowered accuracy, even though the later models have more trees or a different algorithm and more features. That drop is what a leakage-prone split would be expected to hide. PLI remains perfect or nearly perfect once a 50 Hz feature, or a spectrogram, can see the tone. BW versus MA is the unsolved part of the classical problem. Lengthening the window from 2 s to 4 s and changing the model from the tuned forest (0.7766) to gradient boosting (0.8005) coincides with a gain of about two points, but the patients, the feature set and the SNR mix also changed, so the gain is not an isolated ablation.")

    H(doc, "5.2 Exploratory clean-versus-noisy demonstration", 2)
    P(doc, "random_forest_figures/ is Group B. The script cuts record 100 and record 118e_6 into non-overlapping 4 s windows: 451 windows each, 902 in total. A 100-tree forest with class_weight balanced is trained on a stratified 30% split (271 test windows). The saved report gives accuracy 1.0000, and predictions.csv has no disagreements. The script itself warns that windows from the same record fall on both sides of the split and that the score does not estimate generalisation.")
    P(doc, "The score is also easy to understand from the features. Electrode motion at −6 dB inflates RMS from 0.189 to 2.02 and variance from 0.036 to 8.07. A threshold on either feature would already separate these two files. Nothing in this section measures BW against MA against PLI.")
    figure(doc, ROOT / "random_forest_figures" / "confusion_matrix.png",
           "Exploratory two-class confusion matrix (Clean versus Noisy). Accuracy on 271 windows is 1.00. This is not the three-class evaluation.",
           "5. Results", "exploratory_group_B", "exploratory_visualization", width=4.6)
    figure(doc, ROOT / "random_forest_figures" / "feature_importance.png",
           "Exploratory forest importance on the six features defined in random_forest_visualizations.py. The ranking applies only to record 100 versus record 118e_6.",
           "5. Results", "exploratory_group_B", "exploratory_visualization", width=5.6)

    H(doc, "5.3 Saved CNN and ResNet-101 evaluations", 2)
    P(doc, "Three deep-learning evaluations were saved as text reports and confusion matrices. Their training scripts are not in the tree. The test sets are class-balanced, so accuracy and balanced accuracy match. Per-class precision, recall and F1 below are the two-decimal figures printed in the reports. The integer counts are read from the raw matrices and agree with those rounded scores.")
    table(doc,
          ["Run", "Test windows", "Accuracy", "BW P / R / F1", "MA P / R / F1", "PLI P / R / F1"],
          [
              ["CNN, folder cnn_6hours", "1,344 (448 per class)", "0.8772", "0.87 / 0.74 / 0.80", "0.77 / 0.89 / 0.83", "1.00 / 1.00 / 1.00"],
              ["ResNet-101, 6 h folder", "1,344 (448 per class)", "0.8765", "0.77 / 0.89 / 0.83", "0.87 / 0.74 / 0.80", "1.00 / 1.00 / 1.00"],
              ["ResNet-101, 10 h folder", "2,256 (752 per class)", "0.8590", "0.76 / 0.85 / 0.80", "0.83 / 0.73 / 0.78", "1.00 / 1.00 / 1.00"],
          ],
          font=8)
    caption(doc, "Table", "Saved spectrogram-classifier reports. Epochs, learning rates, image size and pretrained weights are not in these files.",
            "5. Results", "cnn_and_resnet101", "original_experiment")
    figure(doc, ROOT / "outputs/milestone3_results/cnn_6hours/cnn_confusion_raw.png",
           "CNN six-hour run, raw confusion matrix. Of 448 BW windows, 332 are correct and 116 are called MA. Of 448 MA windows, 399 are correct and 49 are called BW. All 448 PLI windows are correct.",
           "5. Results", "cnn_6hours", "original_experiment", width=4.8)
    figure(doc, ROOT / "outputs/milestone3_results/resnet101_6hours/resnet_confusion_raw.png",
           "ResNet-101 six-hour run. BW is now the higher-recall class (400/448) and MA is lower (331/448). One PLI window is called MA. Accuracy 0.8765 is within 0.1 point of the CNN on the same test size.",
           "5. Results", "resnet101_6hours", "original_experiment", width=4.8)
    figure(doc, ROOT / "outputs/milestone3_results/resnet101_10hours/resnet_confusion_raw.png",
           "ResNet-101 ten-hour run on 2,256 windows. Accuracy 0.8590. BW 641 correct and 111 called MA; MA 548 correct and 204 called BW; PLI 749 correct, with 1 called BW and 2 called MA.",
           "5. Results", "resnet101_10hours", "original_experiment", width=4.8)
    P(doc, "On a test of 1,344 windows the CNN and ResNet-101 differ by 0.0007 in accuracy. They swap which of BW or MA they prefer, and neither resolves that pair. The ten-hour ResNet-101 test is larger and the accuracy is lower (0.8590), so the longer folder name is not evidence of a better model. The test size matches the final 4 s classical test (2,256), but the ResNet report does not list patient identifiers, so the two test sets are not known to be the same patients. Relative to the classical 0.8005 on a test of that size, the saved networks are higher. The comparison is informative and not controlled: the network inputs, the training recipe and possibly the patients differ, and each saved network appears to have been trained once.")

    H(doc, "5.4 Scores that exist only in the earlier document", 2)
    P(doc, "Complete_Project_Report.pdf evaluates three ResNet-50 heads on a claimed test of 4,500 windows (1,500 per class) and places the improved 2 s forest beside them as “about 77.7%”. The forest figure matches Section 5.1. The network figures do not match any file under outputs/.")
    table(doc,
          ["ID in that document", "Description given there", "Test accuracy given there", "In this repository?"],
          [
              ["—", "Random forest, six features, record-level split", "about 77.7%", "Yes. Exact value 0.7766, but that model has 10 features, not six."],
              ["A", "Baseline ResNet-50, 512-unit head, dropout 0.5", "89.67%", "No saved report or weights."],
              ["B", "A plus L2 of 1e−4 and label smoothing 0.1", "90.87%", "No saved report. The source uses a 128-unit head and dropout 0.4."],
              ["C", "64-unit bottleneck, dropout 0.5", "89.80%", "No saved report or source for this head."],
          ],
          font=8)
    caption(doc, "Table", "Claims in Complete_Project_Report.pdf that could not be reopened as experiment files. They are not used as the primary results.",
            "5. Results", "Complete_Project_Report.pdf", "unverified_document_claim")
    P(doc, "The document’s own discussion says that each configuration was trained once, that validation rose and then overfit on configuration A, and that configuration B was the highest of the three by 1.20 points over A. Those sentences describe runs that are not in the tree. They are not repeated here as findings. The document also fills several method fields with “to be filled”, including window length, STFT length and hop, image size, optimiser and batch size. The current ResNet-50 source does specify 4 s, nperseg 64, hop 16, 128 × 128, Adam and batch 64. Specifying them in code is not the same thing as having the accuracy file.")

    H(doc, "5.5 Filtering figure from the earlier document", 2)
    P(doc, "The earlier document includes one before-and-after figure for the three routed filters and describes it qualitatively. BW input drift is described there as roughly 2 mV, with a large dc spectral line; after the high-pass the baseline is recentred. MA is described as the least successful case, with a suppressed tail above 35 Hz and residual jitter below the cutoff. PLI is described as a narrow 50 Hz line removed by the notch, with the QRS left in place. The document states that RMSE, SNR improvement and a fixed-filter baseline were not computed. This report does not add those numbers.")
    figure(doc, FILTERS,
           "Time- and frequency-domain examples reproduced from Complete_Project_Report.pdf. Subplot titles in the original say that each filter “removes” the noise. The surrounding text of that document is more careful: MA inside the ECG band remains, and no quantitative restoration score was calculated. The plotting script is not in this repository.",
           "5. Results", "prior_project_report", "prior_report_figure", width=6.3)

    H(doc, "5.6 Comparison", 2)
    P(doc, "Across every verified three-class matrix, PLI is the separable class and BW/MA is the residual problem. Scalar 50 Hz power is enough for PLI on these synthetic tones. It is not a test of mains frequency error, harmonics or a wandering line frequency. Spectrograms raise accuracy relative to the best saved classical model when the comparison is allowed to be imperfect: 0.877 and 0.876 on 1,344 windows, and 0.859 on 2,256 windows, against 0.8005 for gradient boosting on 2,256 windows. The exploratory accuracy of 1.00 is left out of that sentence on purpose.")
    P(doc, "The SNR trend is consistent in the two experiments that plotted it. Classification is best when the noise is strongest. A system that only has to work at 0 dB will look better than one that must also work at 15 dB, where the noise signature is weak and the ECG morphology dominates the features.")

    # ---------- 6 ----------
    H(doc, "6. Discussion")
    P(doc, "The project does what the classical half of the design requires. It builds a label by construction, holds out records, extracts features whose 50 Hz member has a physical meaning, and shows where the classifier fails. The failure mode is stable from 54 test windows to 2,256. Calling a BW window MA, or the reverse, would select the wrong filter: a high-pass where a low-pass was required, or the opposite. That is a more serious error than a small drop in accuracy, because the two filters cut opposite ends of the spectrum. PLI errors are rare in the saved matrices, so the notch would seldom be applied to a window that was not tonal, and a tonal window would seldom receive a Butterworth cutoff instead.")
    P(doc, "The dead low-frequency feature is a methodological warning. A band of 0.05–0.5 Hz is the right place to look for baseline wander, but a 2 s Welch spectrum barely samples it, and even the 4 s six-feature model assigned it importance 0. Later features (dominant frequency, zero-crossing rate, mean absolute difference, Hjorth mobility) were added because amplitude and rate-of-change carry BW-versus-MA information that the narrow band did not. The 4 s importance figure ranks those dynamic features above RMS and kurtosis. Whether they close the BW/MA gap is answered by the confusion matrix, not by the ranking: they do not close it.")
    P(doc, "Deep models were saved only as end-point scores. Without the training script, the split list or a learning curve, a higher accuracy cannot be audited. Using the test patients as the validation set, which is what main_cnn.py does, would inflate a reported test score if that code path had been the one that produced the saved matrices. It is not known whether the saved CNN and ResNet-101 runs did this. The limitation is recorded because the only network script in the tree does it.")
    P(doc, "Restoration is the unfinished half. A block diagram and one illustrative figure exist. A function that takes a predicted label and returns a filtered window does not. Until that function exists, “dynamic DSP filtering” is a specification plus an example, not a measured stage of the pipeline. Visual inspection of one window per class cannot support a claim of noise removal. Muscle artifact that shares the QRS band cannot be removed by a 35 Hz low-pass without also taking part of the QRS, and the earlier document already says the residual jitter remains.")
    P(doc, "Generalisation is limited by the way the labels were made. One synthetic or recorded noise is added to a clean excerpt. Real recordings contain mixtures, non-stationary SNR, 60 Hz mains in some countries, and electrode motion, which this training set never asks the model to name. Channel 0 is not guaranteed to be MLII on every record. Windows overlap, so neighbouring excerpts are not independent, even when the record split blocks the worst form of leakage. Each model was trained once. Differences of one accuracy point are not a statistical test.")
    P(doc, "Nothing in the repository supports a clinical claim. The labels are the noise that the program added. No cardiologist annotated the test set. No downstream QRS detector was scored before and after filtering.")

    # ---------- 7 ----------
    H(doc, "7. Conclusion")
    P(doc, "HEARTLINE studies ECG noise classification as the control step for a later filter bank. On public 360 Hz data, with one injected class per window, a record-level 4 s experiment reaches a test accuracy of 0.8005 with histogram gradient boosting on 14 features (2,256 test windows). A tuned random forest on 2 s windows reaches 0.7766 (1,128 test windows). Earlier scores of 0.8333 and 0.8500 come from much smaller tests and are not the headline result. In all four classical matrices, PLI is recognised and BW is confused with MA.")
    P(doc, "Saved spectrogram classifiers score 0.8772 (CNN, 1,344 windows), 0.8765 (ResNet-101, 1,344 windows) and 0.8590 (ResNet-101, 2,256 windows). They make the same BW/MA errors. A ResNet-50 pipeline is written in the source, including the STFT, the 128 × 128 three-channel image, ImageNet initialisation, L2 regularisation and label smoothing, and its output directory is absent. Accuracies above 89% appear only in an earlier PDF and are not treated as repository results.")
    P(doc, "The filter map is a 0.5 Hz high-pass for BW, a 35 Hz low-pass for MA and a 50 Hz notch for PLI, as stated in that PDF. The current code does not apply those filters, and no SNR improvement was measured. An exploratory clean-versus-noisy forest scores 1.00 on records 100 and 118e_6 because those two files are easy to separate, including by RMS alone. That demonstration is not the three-class system.")

    # ---------- 8 ----------
    H(doc, "8. Future Work")
    bullets(doc, [
        "Implement the three filters in code, using the stated order, cutoffs and Q, and record whether filtering is forward-only or forward–backward.",
        "Score restoration on the full test set with SNR improvement, time-domain error against the clean window, and a simple measure of QRS preservation. Compare the routed bank with each filter used on every window.",
        "Keep a patient-level validation split that is not the test split, and save the training script, weights and learning curve with every network run.",
        "Repeat training enough times to put an interval on differences of about one accuracy point.",
        "Add mixture labels and electrode-motion as its own class, instead of illustrating electrode motion only in a two-class figure.",
        "Revisit the 0.05–0.5 Hz feature at both window lengths so that a baseline measure is actually usable by the forest.",
    ])

    # ---------- 9 ----------
    H(doc, "9. References")
    P(doc, "The list follows Complete_Project_Report.pdf. Where that document left a title, initial or venue incomplete, the gap is left open and marked. Entries were not completed from memory.", size=11, italic=True, align="left")
    refs = [
        "[1] Pravin and Ojha, 2022, neural-network-based selection of the denoising filter for noisy ECG, arXiv:2202.00579.",
        "[2] Caldas et al., 2023, detection of wandering baseline and AC interference in ECG with a 1D CNN, Springer, Studies in Computational Intelligence. Full title not stated in the source document.",
        "[3] “Transfer learning in ECG diagnosis: Is it effective?”, 2025, PMC12088039. Author list not stated in the source document.",
        "[4] Bammara and Mousselmal, 2025, ECG classification using 1D–2D transformation and CNN, M2 thesis, University of Bordj Bou Arréridj.",
        "[5] Self-attention-driven ECG denoising, Signals, vol. 6, no. 2, article 26, 2025. Author list not stated in the source document.",
        "[6] Satija, Ramkumar and Manikandan, 2018, IEEE Reviews in Biomedical Engineering, DOI: 10.1109/RBME.2018.2810957.",
        "[7] Li, Rajagopalan and Clifford, 2014, Computer Methods and Programs in Biomedicine, DOI: 10.1016/j.cmpb.2014.09.002.",
        "[8] Chen et al., 2019, IEEE Access, DOI: 10.1109/ACCESS.2019.2944027. Full author list and title not stated in the source document.",
        "[9] Romero et al., baseline wander removal, arXiv:1807.11359. Full title not stated in the source document.",
        "[10] Moody, Muldrow and Mark, 1984, A noise stress test for arrhythmia detectors.",
        "[11] Goldberger et al., 2000, PhysioBank / PhysioNet, DOI: 10.1161/01.CIR.101.23.e215.",
        "[12] MIT-BIH Arrhythmia Database, physionet.org/content/mitdb/1.0.0/; MIT-BIH Noise Stress Test Database, physionet.org/content/nstdb/1.0.0/.",
    ]
    for ref in refs:
        p = doc.add_paragraph()
        _p_fmt(p, before=1, after=4, align="left")
        add_text(p, ref, size=11)

    # ---------- 10 ----------
    H(doc, "10. Appendix")
    H(doc, "10.1 Record lists for the two largest classical sets", 2)
    P(doc, "Final 4 s training records, from dataset_summary.txt: 109, 116, 117, 118, 119, 123, 200, 201, 202, 203, 205, 207, 208, 212, 214, 219, 220, 221, 222, 230, 231, 234. Test records: 104, 105, 107, 122, 209, 210, 228, 232.", align="left")
    P(doc, "Improved 2 s training records: 104, 109, 118, 119, 201, 207, 208, 212, 214, 219, 220, 230, 231, 232, 234. Test records: 105, 107, 122, 123, 209.", align="left")
    H(doc, "10.2 Discrepancies left unresolved", 2)
    bullets(doc, [
        "The 4 s model filename contains “random_forest”, but the object is HistGradientBoostingClassifier and the text report says Gradient Boosting.",
        "The 4 s importance figure is titled as a random forest. The selected booster does not store those importances.",
        "Improved 2 s accuracy is 0.7766 on 10 features. The earlier PDF says “about 77.7%” and “six features”.",
        "ResNet-50 accuracies 89.67%, 90.87% and 89.80% are in the PDF only. The source head (128 units, dropout 0.4, 15 fine-tuning epochs) does not match the PDF head (512 units, dropout 0.5, 45 fine-tuning epochs).",
        "CNN and ResNet-101 training configuration files are missing. CRNN, resnet50_50ep and resnet50_opt2_opt3 directories are missing.",
        "Filter code is missing. The verification figure cannot be regenerated from this tree.",
        "Individual contributions of the four students are not recorded, so none are assigned.",
        "A session year is printed on the formatting sample (a different course). It is not printed on the ECG project document, so it is omitted here.",
    ])
    H(doc, "10.3 Exploratory feature means", 2)
    P(doc, "Means over 451 windows, from random_forest_figures/extracted_features.csv, using that script’s bands rather than the three-class bands.", align="left")
    table(doc,
          ["Feature", "Record 100 (Clean)", "Record 118e_6 (Noisy)"],
          [
              ["RMS", "0.1889", "2.0213"],
              ["Variance", "0.0360", "8.0733"],
              ["Kurtosis", "26.68", "2.76"],
              ["Low-frequency power (0.1–5 Hz)", "0.00594", "5.032"],
              ["50 Hz power (48–52 Hz)", "0.000061", "0.000879"],
              ["High-frequency power (20–100 Hz)", "0.00749", "0.0383"],
          ],
          font=10)
    caption(doc, "Table", "Exploratory feature means. Group B only.", "10. Appendix", "exploratory_group_B", "exploratory_visualization")

    out_docx = OUT / "DSP_Project_Report.docx"
    doc.save(out_docx)
    print("saved", out_docx)
    return FIGS


if __name__ == "__main__":
    figs = build()
    import csv
    inv = OUT / "figure_inventory.csv"
    with inv.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["number", "source_path", "caption", "experiment_id", "chapter", "origin"])
        w.writeheader()
        w.writerows(figs)
    print("figures", len(figs), "->", inv)
